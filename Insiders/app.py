from flask import Flask, jsonify, request
from flask_cors import CORS
import openpyxl
import os
import time
from datetime import datetime
import requests as req_lib

from fraud_model_lr import predict_lr
from fraud_model_xgb import predict_xgb

app = Flask(__name__)
CORS(app)

USER = {
    "name":    "Ravi Kumar",
    "balance": 100000.0,
    "pin":     "1234",
    "cc_num":  3573030041201292,
}

BASE_DIR     = os.path.dirname(os.path.abspath(__file__))
HISTORY_DIR  = os.path.join(BASE_DIR, "History")
HISTORY_FILE = os.path.join(HISTORY_DIR, "History.xlsx")
SUMMARY_FILE = os.path.join(HISTORY_DIR, "Summary.xlsx")

os.makedirs(HISTORY_DIR, exist_ok=True)

_merchant_coord_cache = {}

_session_txns = []   # list of {"unix_time": float, "amount": float}



def rule_based_fraud_check(amount, balance):
    """
    Returns (is_fraud: bool, decision: str, flags: dict)

    Rules applied in order (first match wins):
      R1. amount >= 90% of current balance              -> FRAUD_ALERT
      R2. amount >= 70% of current balance              -> FRAUD_ALERT
      R3. amount >= 5x session average (3+ prior txns)  -> FRAUD_ALERT
      R4. 3+ txns in the last 60 seconds                -> LOGOUT_BLOCK
    """
    now = time.time()

    flags = {
        "balance_wipeout":         False,
        "large_pct_of_balance":    False,
        "abnormal_vs_session_avg": False,
        "rapid_txns":              False,
    }

    if balance > 0 and (amount / balance) >= 0.90:
        flags["balance_wipeout"] = True
        return True, "FRAUD_ALERT", flags

    if balance > 0 and (amount / balance) >= 0.70:
        flags["large_pct_of_balance"] = True
        return True, "FRAUD_ALERT", flags

    if len(_session_txns) >= 3:
        avg = sum(t["amount"] for t in _session_txns) / len(_session_txns)
        if avg > 0 and amount >= 5 * avg:
            flags["abnormal_vs_session_avg"] = True
            return True, "FRAUD_ALERT", flags

    recent_60s = [t for t in _session_txns if now - t["unix_time"] <= 60]
    if len(recent_60s) >= 3:
        flags["rapid_txns"] = True
        return True, "LOGOUT_BLOCK", flags

    return False, "PASS", flags


def build_rule_response(amount, balance, rule_decision, rule_flags):
    """Build the JSON payload for a rule-engine fraud block."""
    is_logout = rule_decision == "LOGOUT_BLOCK"
    pct = round((amount / balance * 100) if balance > 0 else 100, 1)

    sus_flags = {
        "time_sus":                0,
        "location_sus":            0,
        "amount_sus":              int(
            rule_flags.get("balance_wipeout") or
            rule_flags.get("large_pct_of_balance") or
            rule_flags.get("abnormal_vs_session_avg")
        ),
        "merchant_sus":            0,
        "sus_score":               1,
        "balance_wipeout":         int(rule_flags.get("balance_wipeout", False)),
        "large_pct_of_balance":    int(rule_flags.get("large_pct_of_balance", False)),
        "abnormal_vs_session_avg": int(rule_flags.get("abnormal_vs_session_avg", False)),
        "rapid_txns":              int(rule_flags.get("rapid_txns", False)),
        "balance_at_check":        round(balance, 2),
        "attempted_amount":        round(amount, 2),
        "pct_of_balance":          pct,
    }

    if is_logout:
        msg = "Transaction blocked: Rapid-fire transactions detected. Account locked for your protection."
    elif rule_flags.get("balance_wipeout"):
        msg = f"Transaction blocked: Attempting to spend {pct}% of your balance (Rs.{amount:,.2f} of Rs.{balance:,.2f}) in one transaction."
    elif rule_flags.get("large_pct_of_balance"):
        msg = f"Transaction blocked: Attempting to spend {pct}% of your balance in a single transaction is flagged as suspicious."
    else:
        msg = "Transaction blocked: Amount is abnormally high compared to your recent session spending."

    payload = {
        "fraud_blocked":     True,
        "final_decision":    rule_decision,
        "fraud_probability": 0.99,
        "sus_flags":         sus_flags,
        "model_used":        "RULE_ENGINE",
        "xgb":               None,
        "error":             msg,
    }
    if is_logout:
        payload["force_logout"] = True
    return payload



# Default coords (Hyderabad, India) — used when geo lookup fails
DEFAULT_LAT  = 17.3850
DEFAULT_LON  = 78.4867
DEFAULT_CITY = "Hyderabad"
DEFAULT_LOC  = "Hyderabad, Telangana, IN"


def get_geo():
    """Returns (lat, lon, city, location_str) — NEVER returns None."""
    try:
        from dotenv import load_dotenv
        load_dotenv()
        token = os.getenv("IPINFO_TOKEN", "")
        res = req_lib.get(f"https://ipinfo.io/json?token={token}", timeout=5)
        geo = res.json()
        loc = geo.get("loc", "")
        city = geo.get("city", DEFAULT_CITY)
        region = geo.get("region", "")
        country = geo.get("country", "")
        location_str = f"{city}, {region}, {country}"
        if loc:
            lat, lon = map(float, loc.split(","))
            return lat, lon, city, location_str
    except Exception as e:
        print(f"[GEO] lookup failed ({e}), using default location")
    return DEFAULT_LAT, DEFAULT_LON, DEFAULT_CITY, DEFAULT_LOC


def geocode_merchant(merchant, city, fallback_lat, fallback_lon):
    cache_key = f"{merchant.lower().strip()}|{city.lower().strip()}"
    if cache_key in _merchant_coord_cache:
        return _merchant_coord_cache[cache_key]

    headers = {"User-Agent": "NexusBank-FraudDetection/1.0"}
    for q in [f"{merchant}, {city}", merchant]:
        try:
            resp = req_lib.get(
                "https://nominatim.openstreetmap.org/search",
                params={"q": q, "format": "json", "limit": 1},
                headers=headers, timeout=6,
            )
            results = resp.json()
            if results:
                lat = float(results[0]["lat"])
                lon = float(results[0]["lon"])
                _merchant_coord_cache[cache_key] = (lat, lon)
                return lat, lon
        except Exception:
            pass
        time.sleep(1)

    _merchant_coord_cache[cache_key] = (fallback_lat, fallback_lon)
    return fallback_lat, fallback_lon


#  ML FRAUD PIPELINE

def run_fraud_pipeline(amount, merchant, category, session_time_hms, now, lat, lon, city):
    merch_lat, merch_lon = geocode_merchant(merchant, city, lat, lon)

    fraud_input = {
        "cc_num":                USER["cc_num"],
        "amt":                   float(amount),
        "lat":                   lat,
        "long":                  lon,
        "merch_lat":             merch_lat,
        "merch_long":            merch_lon,
        "category":              category,
        "merchant":              merchant,
        "city":                  city,
        "session_time":          session_time_hms,
        "trans_date_trans_time": now.strftime("%d/%m/%Y %H:%M"),
        "unix_time":             int(time.time()),
    }

    # Logistic Regression (first gate)
    try:
        lr_result = predict_lr(fraud_input)
    except Exception as e:
        print(f"[LR ERROR] {e} — ML unavailable, authenticating")
        return {
            "final_decision":    "AUTHENTICATE",
            "model_used":        "LR_UNAVAILABLE",
            "fraud_probability": 0.0,
            "sus_flags":         {},
            "xgb":               None,
        }

    if lr_result["decision"] == "AUTHENTICATE":
        return {
            "final_decision":    "AUTHENTICATE",
            "model_used":        "LR",
            "fraud_probability": lr_result["fraud_probability"],
            "sus_flags":         lr_result["sus_flags"],
            "xgb":               None,
        }

    if lr_result["decision"] == "FRAUD_ALERT":
        return {
            "final_decision":    "FRAUD_ALERT",
            "model_used":        "LR",
            "fraud_probability": lr_result["fraud_probability"],
            "sus_flags":         lr_result["sus_flags"],
            "xgb":               None,
        }

    # XGBoost (second gate)
    try:
        xgb_result = predict_xgb(fraud_input)
    except Exception as e:
        print(f"[XGB ERROR] {e} — XGB unavailable, authenticating")
        return {
            "final_decision":    "AUTHENTICATE",
            "model_used":        "XGB_UNAVAILABLE",
            "fraud_probability": lr_result["fraud_probability"],
            "sus_flags":         lr_result["sus_flags"],
            "xgb":               None,
        }

    return {
        "final_decision":    xgb_result["decision"],
        "model_used":        "XGBoost",
        "fraud_probability": lr_result["fraud_probability"],
        "sus_flags":         lr_result["sus_flags"],
        "xgb":               xgb_result,
    }


#  EXCEL HELPERS

def save_history_excel(all_txns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Date", "Time", "Merchant", "Category", "Txn_Type",
                "Amount", "Net Debit", "Location", "Session_Time"])
    for t in all_txns:
        ws.append([
            t.get("date"), t.get("time"), t.get("merchant"),
            t.get("category"), t.get("txn_type") or t.get("type"),
            t.get("amount"), t.get("net_debit") or t.get("net debit"),
            t.get("location"), t.get("session_time", ""),
        ])
    wb.save(HISTORY_FILE)


def load_history_excel():
    if not os.path.exists(HISTORY_FILE):
        return []
    try:
        wb = openpyxl.load_workbook(HISTORY_FILE)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return []
        headers = [str(h).lower().replace(" ", "_") for h in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]
    except Exception:
        return []


def save_summary_excel(last_txn, stats, merchant_averages):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Last Transaction"
    ws1.append(["Field", "Value"])
    for key, val in last_txn.items():
        ws1.append([key, val])
    ws2 = wb.create_sheet("Stats")
    ws2.append(["Metric", "Value"])
    for key, val in stats.items():
        ws2.append([key, val])
    ws3 = wb.create_sheet("Merchant Averages")
    ws3.append(["Merchant", "Avg Amount", "Txn Count"])
    for m, info in merchant_averages.items():
        ws3.append([info["merchant"], info["avg_amount"], info["count"]])
    wb.save(SUMMARY_FILE)


def load_summary_excel():
    if not os.path.exists(SUMMARY_FILE):
        return None, None, None
    try:
        wb = openpyxl.load_workbook(SUMMARY_FILE)
        if "Last Transaction" not in wb.sheetnames or "Stats" not in wb.sheetnames:
            return None, None, None
        ws1 = wb["Last Transaction"]
        last_txn = {row[0]: row[1] for row in list(ws1.iter_rows(values_only=True))[1:] if row[0]}
        ws2 = wb["Stats"]
        stats = {row[0]: row[1] for row in list(ws2.iter_rows(values_only=True))[1:] if row[0]}
        merchant_averages = {}
        if "Merchant Averages" in wb.sheetnames:
            ws3 = wb["Merchant Averages"]
            for row in list(ws3.iter_rows(values_only=True))[1:]:
                if row[0]:
                    merchant_averages[str(row[0]).lower()] = {
                        "merchant": row[0], "avg_amount": row[1], "count": row[2]
                    }
        if not last_txn or not stats:
            return None, None, None
        return last_txn, stats, merchant_averages
    except Exception:
        return None, None, None


def compute_and_save(session_txns):
    all_txns = load_history_excel()
    all_txns.extend(session_txns)
    save_history_excel(all_txns)
    last = session_txns[-1]
    last_txn = {k: last[k] for k in ("date", "time", "merchant", "amount", "location")}
    total_txns   = len(all_txns)
    total_amount = sum(float(t["amount"]) for t in all_txns)
    avg_amount   = round(total_amount / total_txns, 2)

    def parse_date(val):
        if isinstance(val, datetime):
            return val
        return datetime.strptime(str(val), "%Y-%m-%d")

    dates     = [parse_date(t["date"]) for t in all_txns]
    day_span  = max((max(dates) - min(dates)).days + 1, 1)
    week_span = max(day_span / 7, 1)
    stats = {
        "total_txns":         total_txns,
        "avg_amount_per_txn": avg_amount,
        "avg_txns_per_day":   round(total_txns / day_span, 2),
        "avg_txns_per_week":  round(total_txns / week_span, 2),
    }
    merch_data = {}
    for t in all_txns:
        key = str(t["merchant"]).lower()
        if key not in merch_data:
            merch_data[key] = {"merchant": t["merchant"], "total": 0.0, "count": 0}
        merch_data[key]["total"] = round(merch_data[key]["total"] + float(t["amount"]), 2)
        merch_data[key]["count"] += 1
    merchant_averages = {
        key: {
            "merchant":   v["merchant"],
            "avg_amount": round(v["total"] / v["count"], 2),
            "count":      v["count"],
        }
        for key, v in merch_data.items()
    }
    save_summary_excel(last_txn, stats, merchant_averages)
    return all_txns, stats, merchant_averages



@app.route("/api/user", methods=["GET"])
def get_user():
    return jsonify({"name": USER["name"]})


@app.route("/api/balance", methods=["POST"])
def get_balance():
    data = request.json or {}
    if data.get("pin") != USER["pin"]:
        return jsonify({"error": "Incorrect PIN"}), 401
    return jsonify({"balance": USER["balance"], "name": USER["name"]})


@app.route("/api/deposit", methods=["POST"])
def deposit():
    data = request.json or {}
    if data.get("pin") != USER["pin"]:
        return jsonify({"error": "Incorrect PIN"}), 401
    amount = float(data.get("amount", 0))
    if amount <= 0:
        return jsonify({"error": "Amount must be greater than zero"}), 400
    USER["balance"] = round(USER["balance"] + amount, 2)
    return jsonify({"balance": USER["balance"], "deposited": amount})


@app.route("/api/transaction", methods=["POST"])
def transaction():
    server_received_ts = time.time() * 1000
    data = request.json or {}

    # Auth check
    if data.get("pin") != USER["pin"]:
        return jsonify({"error": "Incorrect PIN"}), 401

    amount   = float(data.get("amount", 0))
    merchant = str(data.get("merchant", "")).strip()
    category = data.get("category", "Others")
    txn_type = data.get("txn_type", "UPI")

    if amount <= 0:
        return jsonify({"error": "Amount must be greater than zero"}), 400
    if not merchant:
        return jsonify({"error": "Merchant name cannot be empty"}), 400
    if amount > USER["balance"]:
        return jsonify({"error": "Insufficient balance"}), 400

    is_rule_fraud, rule_decision, rule_flags = rule_based_fraud_check(amount, USER["balance"])
    print(f"[RULE CHECK] amount={amount}, balance={USER['balance']}, fraud={is_rule_fraud}, decision={rule_decision}, flags={rule_flags}")

    if is_rule_fraud:
        payload = build_rule_response(amount, USER["balance"], rule_decision, rule_flags)
        print(f"[RULE BLOCK] {payload['error']}")
        return jsonify(payload), 403

    # ML FRAUD PIPELINE
    session_start_ts = data.get("session_start_ts")
    if session_start_ts:
        session_elapsed_s = (server_received_ts - float(session_start_ts)) / 1000
        total_secs = int(session_elapsed_s)
        session_time_hms = f"{total_secs//3600:02d}:{(total_secs%3600)//60:02d}:{total_secs%60:02d}"
        mins = int(session_elapsed_s // 60)
        secs = session_elapsed_s % 60
        session_time_str = f"{mins}m {secs:.2f}s" if mins else f"{secs:.2f}s"
    else:
        session_time_hms = "00:00:00"
        session_time_str = "Unknown"

    # get_geo() now NEVER returns None
    lat, lon, city, location = get_geo()
    now = datetime.now()

    try:
        fraud_result = run_fraud_pipeline(
            amount, merchant, category, session_time_hms, now, lat, lon, city
        )
    except Exception as e:
        print(f"[ML PIPELINE ERROR] {e} — approving transaction, rule engine already passed")
        fraud_result = {
            "final_decision":    "AUTHENTICATE",
            "model_used":        "PIPELINE_ERROR",
            "fraud_probability": 0.0,
            "sus_flags":         {},
            "xgb":               None,
        }

    final_decision = fraud_result["final_decision"]
    print(f"[ML RESULT] decision={final_decision}, prob={fraud_result.get('fraud_probability')}")

    if final_decision in ("FRAUD_ALERT", "LOGOUT_BLOCK"):
        payload = {
            "fraud_blocked":     True,
            "final_decision":    final_decision,
            "fraud_probability": fraud_result["fraud_probability"],
            "sus_flags":         fraud_result["sus_flags"],
            "model_used":        fraud_result["model_used"],
            "xgb":               fraud_result["xgb"],
            "error": (
                "Transaction blocked: High fraud probability detected by ML model."
                if final_decision == "FRAUD_ALERT"
                else "Transaction blocked: Suspicious behavioral pattern. Account locked for safety."
            ),
        }
        if final_decision == "LOGOUT_BLOCK":
            payload["force_logout"] = True
        return jsonify(payload), 403

    USER["balance"] = round(USER["balance"] - amount, 2)
    _session_txns.append({"unix_time": time.time(), "amount": amount})

    txn = {
        "date":         now.strftime("%Y-%m-%d"),
        "time":         now.strftime("%H:%M:%S"),
        "merchant":     merchant,
        "category":     category,
        "txn_type":     txn_type,
        "amount":       amount,
        "net_debit":    amount,
        "location":     location,
        "session_time": session_time_str,
    }

    return jsonify({
        "transaction":  txn,
        "balance":      USER["balance"],
        "fraud_result": fraud_result,
    })


@app.route("/predict", methods=["POST"])
def predict():
    data = request.get_json()
    required = [
        "cc_num", "amt", "lat", "long", "merch_lat", "merch_long",
        "category", "merchant", "city", "session_time",
        "trans_date_trans_time", "unix_time"
    ]
    for field in required:
        if field not in data:
            return jsonify({"error": f"Missing field: {field}"}), 400

    try:
        lr_result = predict_lr(data)
    except Exception as e:
        return jsonify({"error": f"LR model error: {e}"}), 500

    if lr_result["decision"] in ("AUTHENTICATE", "FRAUD_ALERT"):
        return jsonify({
            "final_decision":    lr_result["decision"],
            "model_used":        "LR",
            "fraud_probability": lr_result["fraud_probability"],
            "sus_flags":         lr_result["sus_flags"],
            "xgb":               None,
        })

    try:
        xgb_result = predict_xgb(data)
    except Exception as e:
        return jsonify({"error": f"XGB model error: {e}"}), 500

    return jsonify({
        "final_decision":    xgb_result["decision"],
        "model_used":        "XGBoost",
        "fraud_probability": lr_result["fraud_probability"],
        "sus_flags":         lr_result["sus_flags"],
        "xgb":               xgb_result,
    })


@app.route("/api/history", methods=["GET"])
def history():
    return jsonify({"transactions": load_history_excel()})


@app.route("/api/summary", methods=["GET"])
def summary():
    last, stats, merch = load_summary_excel()
    if not last:
        return jsonify({"available": False})
    return jsonify({"available": True, "last": last, "stats": stats, "merchants": merch})


@app.route("/api/save", methods=["POST"])
def save():
    data = request.json or {}
    session_txns = data.get("transactions", [])
    if not session_txns:
        return jsonify({"message": "No transactions to save"})
    all_txns, stats, merch = compute_and_save(session_txns)
    return jsonify({"message": "Saved successfully", "total": len(all_txns), "stats": stats})


@app.route("/health", methods=["GET"])
def health():
    return jsonify({
        "status":       "ok",
        "balance":      USER["balance"],
        "session_txns": len(_session_txns),
    })


if __name__ == "__main__":
    app.run(debug=True, port=5000)
