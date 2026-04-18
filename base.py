import openpyxl
import os
from datetime import datetime
import requests
from dotenv import load_dotenv


USER = {
    "name"   : "Ravi Kumar",
    "balance": 50000.0,
    "pin"    : "1234"
}

session_transactions = []
SESSION_START = datetime.now()


def format_duration(seconds):
    """Converts seconds into HH:MM:SS string."""
    hours, remainder = divmod(int(seconds), 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


#  File Paths
HISTORY_FILE = r"C:\Users\Roshini\Desktop\JNTUH Hackathon\Insiders-Cloud-Hackathon\History\History.xlsx"
SUMMARY_FILE = r"C:\Users\Roshini\Desktop\JNTUH Hackathon\Insiders-Cloud-Hackathon\History\Summary.xlsx"

TRANSACTION_TYPES = ["Card", "UPI"]
CATEGORIES = [
    "Food & Dining", "Shopping", "Travel", "Entertainment",
    "Utilities", "Healthcare", "Education", "Others"
]


#  Excels

def save_history_excel(all_txns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Date", "Time", "Merchant", "Category",
               "Txn_Type", "Amount", "Net Debit", "Location", "Time per Transaction"])
    for t in all_txns:
        ws.append([
            t.get("date"),
            t.get("time"),
            t.get("merchant"),
            t.get("category"),
            t.get("txn_type") or t.get("type"),
            t.get("amount"),
            t.get("net_debit") or t.get("net debit"),
            t.get("location"),
            t.get("time_per_txn", ""),
        ])
    wb.save(HISTORY_FILE)


def load_history_excel():
    if not os.path.exists(HISTORY_FILE):
        return []
    try:
        wb = openpyxl.load_workbook(HISTORY_FILE)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return []
        headers = [str(h).lower().replace(" ", "_") for h in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]
    except Exception:
        return []


def save_summary_excel(last_txn, stats, merchant_averages):
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Last Transaction"
    ws1.append(["Field", "Value"])
    for key, val in last_txn.items():
        ws1.append([key, val])

    ws2 = wb.create_sheet("Stats")
    ws2.append(["Metric", "Value"])
    for key, val in stats.items():
        ws2.append([key, val])

    ws3 = wb.create_sheet("Merchant Averages")
    ws3.append(["Merchant", "Avg Amount", "Txn Count"])
    for m, info in merchant_averages.items():
        ws3.append([info["merchant"], info["avg_amount"], info["count"]])

    wb.save(SUMMARY_FILE)


def load_summary_excel():
    if not os.path.exists(SUMMARY_FILE):
        return None, None, None
    try:
        wb = openpyxl.load_workbook(SUMMARY_FILE)

        if "Last Transaction" not in wb.sheetnames or "Stats" not in wb.sheetnames:
            return None, None, None

        ws1 = wb["Last Transaction"]
        last_txn = {}
        for row in list(ws1.iter_rows(values_only=True))[1:]:
            if row[0] is not None:
                last_txn[row[0]] = row[1]

        ws2 = wb["Stats"]
        stats = {}
        for row in list(ws2.iter_rows(values_only=True))[1:]:
            if row[0] is not None:
                stats[row[0]] = row[1]

        merchant_averages = {}
        if "Merchant Averages" in wb.sheetnames:
            ws3 = wb["Merchant Averages"]
            for row in list(ws3.iter_rows(values_only=True))[1:]:
                if row[0] is not None:
                    merchant_averages[str(row[0]).lower()] = {
                        "merchant"  : row[0],
                        "avg_amount": row[1],
                        "count"     : row[2],
                    }

        if not last_txn or not stats:
            return None, None, None

        return last_txn, stats, merchant_averages

    except Exception:
        return None, None, None


#  Other

def pick_from_list(prompt, options):
    print(f"\n{prompt}")
    for i, opt in enumerate(options, 1):
        print(f"  {i}. {opt}")
    while True:
        try:
            choice = int(input("Enter number: "))
            if 1 <= choice <= len(options):
                return options[choice - 1]
            print("  Pick a valid number.")
        except ValueError:
            print("  Invalid input, enter a number.")


def verify_pin():
    pin = input("Enter PIN: ").strip()
    if pin != USER["pin"]:
        raise PermissionError("Incorrect PIN.")


# Features

def check_balance():
    print("\n── Check Balance ──")
    try:
        verify_pin()
        print(f"\n  Account Holder : {USER['name']}")
        print(f"  Balance        : Rs.{USER['balance']:,.2f}")
    except PermissionError as pe:
        print("Access Denied:", pe)


def deposit():
    print("\n── Deposit Money ──")
    try:
        verify_pin()
        amount = float(input("Enter deposit amount (Rs.): "))
        if amount <= 0:
            raise ValueError("Amount must be greater than zero.")
        USER["balance"] = round(USER["balance"] + amount, 2)
        print(f"\n  Rs.{amount:,.2f} deposited successfully.")
        print(f"  New Balance: Rs.{USER['balance']:,.2f}")
    except PermissionError as pe:
        print("Access Denied:", pe)
    except ValueError as ve:
        print("Invalid input:", ve)
    except Exception as e:
        print("Something went wrong:", e)


def make_transaction():
    print("\n── Make Transaction ──")
    txn_start = datetime.now()   #stater
    #calculating the session time(subtracting the trandone-traninit)
    try:
        verify_pin()

        amount = float(input("Enter transaction amount (Rs.): "))
        if amount <= 0:
            raise ValueError("Amount must be greater than zero.")
        if amount > USER["balance"]:
            print(f"  Insufficient balance. Available: Rs.{USER['balance']:,.2f}")
            return

        txn_type = pick_from_list("Select Transaction Type:", TRANSACTION_TYPES)
        category = pick_from_list("Select Category:", CATEGORIES)
        merchant = input("Enter Merchant / Sender name: ").strip()
        if not merchant:
            raise ValueError("Merchant name cannot be empty.")

        load_dotenv()
        try:
            token = os.getenv("IPINFO_TOKEN")
            res = requests.get(f"https://ipinfo.io/json?token={token}", timeout=5)
            data = res.json()
            location = f"{data.get('city')}, {data.get('region')}, {data.get('country')}"
        except Exception:
            location = "Unknown"

        txn_end  = datetime.now()   #stopper
        duration = (txn_end - txn_start).total_seconds()

        timestamp = txn_end
        date_str  = timestamp.strftime("%Y-%m-%d")
        time_str  = timestamp.strftime("%H:%M:%S")

        net_debit       = amount
        USER["balance"] = round(USER["balance"] - net_debit, 2)

        txn = {
            "date"        : date_str,
            "time"        : time_str,
            "merchant"    : merchant,
            "category"    : category,
            "txn_type"    : txn_type,
            "amount"      : amount,
            "net_debit"   : net_debit,
            "location"    : location,
            "time_per_txn": format_duration(duration),  # HH:MM:SS taken for this transaction
        }
        session_transactions.append(txn)

        print(f"\n  Transaction Successful!")
        print(f"  Amount Debited   : Rs.{net_debit:.2f}")
        print(f"  Remaining Balance: Rs.{USER['balance']:,.2f}")

    except PermissionError as pe:
        print("Access Denied:", pe)
    except ValueError as ve:
        print("Invalid input:", ve)
    except Exception as e:
        print("Something went wrong:", e)


def view_session_history():
    print("\n── Transactions This Session ──")
    if not session_transactions:
        print("  No transactions made yet this session.")
        return

    print(f"\n{'Date':<12}{'Time':<10}{'Merchant':<20}{'Category':<20}"
          f"{'Type':<6}{'Amount':>10}{'Net':>10}")
    print("-" * 90)
    for t in session_transactions:
        print(f"{t['date']:<12}{t['time']:<10}{t['merchant']:<20}"
              f"{t['category']:<20}{t['txn_type']:<6}"
              f"Rs.{t['amount']:>7.2f}  Rs.{t['net_debit']:>7.2f}")


def view_last_transaction():
    print("\n── Last Transaction (from file) ──")
    last, _, _ = load_summary_excel()
    if not last:
        print("  No transactions on file yet. Make a transaction and Exit & Save first.")
        return
    print(f"  Date     : {last.get('date')}")
    print(f"  Time     : {last.get('time')}")
    print(f"  Merchant : {last.get('merchant')}")
    print(f"  Amount   : Rs.{float(last.get('amount', 0)):.2f}")
    print(f"  Location : {last.get('location')}")


def view_averages():
    print("\n── Averages & Stats (from file) ──")
    _, stats, merchant_averages = load_summary_excel()
    if not stats:
        print("  No stats available yet. Make a transaction and Exit & Save first.")
        return

    print(f"\n  Average amount per transaction : Rs.{float(stats.get('avg_amount_per_txn', 0)):.2f}")
    print(f"  Average transactions per day   : {float(stats.get('avg_txns_per_day', 0)):.2f}")
    print(f"  Average transactions per week  : {float(stats.get('avg_txns_per_week', 0)):.2f}")
    print(f"  Total transactions on record   : {stats.get('total_txns')}")

    if merchant_averages:
        print(f"\n  Average spend per Merchant:")
        print(f"  {'Merchant':<22}{'Avg Amount':>12}{'Txn Count':>12}")
        print("  " + "-" * 46)
        for m, info in sorted(merchant_averages.items(), key=lambda x: -x[1]["avg_amount"]):
            print(f"  {info['merchant']:<22}Rs.{info['avg_amount']:>9.2f}{info['count']:>12}")


#  Save to Files on Exit

def save_on_exit():
    if not session_transactions:
        print("\n  No new transactions to save.")
        return

    all_txns = load_history_excel()
    all_txns.extend(session_transactions)
    save_history_excel(all_txns)

    last = session_transactions[-1]
    last_txn = {
        "date"    : last["date"],
        "time"    : last["time"],
        "merchant": last["merchant"],
        "amount"  : last["amount"],
        "location": last["location"],
    }

    total_txns   = len(all_txns)
    total_amount = sum(float(t["amount"]) for t in all_txns)
    avg_amount   = round(total_amount / total_txns, 2)

    def parse_date(val):
        if isinstance(val, datetime):
            return val
        return datetime.strptime(str(val), "%Y-%m-%d")

    dates     = [parse_date(t["date"]) for t in all_txns]
    earliest  = min(dates)
    latest    = max(dates)
    day_span  = max((latest - earliest).days + 1, 1)
    week_span = max(day_span / 7, 1)

    stats = {
        "total_txns"        : total_txns,
        "avg_amount_per_txn": avg_amount,
        "avg_txns_per_day"  : round(total_txns / day_span, 2),
        "avg_txns_per_week" : round(total_txns / week_span, 2),
    }

    merch_data = {}
    for t in all_txns:
        key = str(t["merchant"]).lower()
        if key not in merch_data:
            merch_data[key] = {"merchant": t["merchant"], "total": 0.0, "count": 0}
        merch_data[key]["total"] = round(merch_data[key]["total"] + float(t["amount"]), 2)
        merch_data[key]["count"] += 1

    merchant_averages = {
        key: {
            "merchant"  : v["merchant"],
            "avg_amount": round(v["total"] / v["count"], 2),
            "count"     : v["count"],
        }
        for key, v in merch_data.items()
    }

    save_summary_excel(last_txn, stats, merchant_averages)

    print(f"\n  Saved successfully!")
    print(f"  -> {HISTORY_FILE}")
    print(f"  -> {SUMMARY_FILE}")


#  Main Menu

def menu():
    print(f"   Welcome, {USER['name']}!")
    
    while True:
        print("\n── Menu ──")
        print("  1. Check Balance")
        print("  2. Deposit Money")
        print("  3. Make a Transaction")
        print("  4. View Session Transactions")
        print("  5. View Last Transaction (File)")
        print("  6. View Averages & Stats (File)")
        print("  7. Exit & Save")

        try:
            choice = int(input("\nEnter your choice: "))
        except ValueError:
            print("  Please enter a valid number.")
            continue

        if choice == 1:
            check_balance()
        elif choice == 2:
            deposit()
        elif choice == 3:
            make_transaction()
        elif choice == 4:
            view_session_history()
        elif choice == 5:
            view_last_transaction()
        elif choice == 6:
            view_averages()
        elif choice == 7:
            save_on_exit()
            print("\n  Goodbye, have a great day!")
            break
        else:
            print("  Invalid choice. Pick 1–7.")


menu()
