from flask import Flask, jsonify, request
from flask_cors import CORS
import openpyxl
import os
from datetime import datetime
import requests as req_lib
import time

app = Flask(__name__)
CORS(app)

# ── State ─────────────────────────────────────────────────────────────────────
USER = {
    "name":    "Ravi Kumar",
    "balance": 50000.0,
    "pin":     "1234",
}

BASE_DIR     = os.path.dirname(os.path.abspath(__file__))
HISTORY_DIR  = os.path.join(BASE_DIR, "History")
HISTORY_FILE = os.path.join(HISTORY_DIR, "History.xlsx")
SUMMARY_FILE = os.path.join(HISTORY_DIR, "Summary.xlsx")

os.makedirs(HISTORY_DIR, exist_ok=True)

# ── Helpers ───────────────────────────────────────────────────────────────────
def format_duration(seconds: float) -> str:
    """Convert seconds to readable string e.g. '1m 23.45s' or '5.32s'"""
    seconds = max(0.0, seconds)
    mins = int(seconds // 60)
    secs = seconds % 60
    if mins:
        return f"{mins}m {secs:.2f}s"
    return f"{secs:.2f}s"

# ── Excel helpers ─────────────────────────────────────────────────────────────
HISTORY_HEADERS = [
    "Date", "Time", "Merchant", "Category", "Txn_Type",
    "Amount", "Net Debit", "Location",
    "Session_Time",       # how long user has been in the app when txn was made
    "Txn_Duration_ms",    # ms from user clicking Pay to server receiving request
]

def save_history_excel(all_txns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(HISTORY_HEADERS)
    for t in all_txns:
        ws.append([
            t.get("date"),
            t.get("time"),
            t.get("merchant"),
            t.get("category"),
            t.get("txn_type") or t.get("type"),
            t.get("amount"),
            t.get("net_debit") or t.get("net debit"),
            t.get("location"),
            t.get("session_time", ""),
            t.get("txn_duration_ms", ""),
        ])
    wb.save(HISTORY_FILE)

def load_history_excel():
    if not os.path.exists(HISTORY_FILE):
        return []
    try:
        wb = openpyxl.load_workbook(HISTORY_FILE)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return []
        headers = [str(h).lower().replace(" ", "_") for h in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]
    except Exception:
        return []

def save_summary_excel(last_txn, stats, merchant_averages):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Last Transaction"
    ws1.append(["Field", "Value"])
    for key, val in last_txn.items():
        ws1.append([key, val])
    ws2 = wb.create_sheet("Stats")
    ws2.append(["Metric", "Value"])
    for key, val in stats.items():
        ws2.append([key, val])
    ws3 = wb.create_sheet("Merchant Averages")
    ws3.append(["Merchant", "Avg Amount", "Txn Count"])
    for m, info in merchant_averages.items():
        ws3.append([info["merchant"], info["avg_amount"], info["count"]])
    wb.save(SUMMARY_FILE)

def load_summary_excel():
    if not os.path.exists(SUMMARY_FILE):
        return None, None, None
    try:
        wb = openpyxl.load_workbook(SUMMARY_FILE)
        if "Last Transaction" not in wb.sheetnames or "Stats" not in wb.sheetnames:
            return None, None, None
        ws1 = wb["Last Transaction"]
        last_txn = {row[0]: row[1] for row in list(ws1.iter_rows(values_only=True))[1:] if row[0]}
        ws2 = wb["Stats"]
        stats = {row[0]: row[1] for row in list(ws2.iter_rows(values_only=True))[1:] if row[0]}
        merchant_averages = {}
        if "Merchant Averages" in wb.sheetnames:
            ws3 = wb["Merchant Averages"]
            for row in list(ws3.iter_rows(values_only=True))[1:]:
                if row[0]:
                    merchant_averages[str(row[0]).lower()] = {
                        "merchant": row[0], "avg_amount": row[1], "count": row[2]
                    }
        if not last_txn or not stats:
            return None, None, None
        return last_txn, stats, merchant_averages
    except Exception:
        return None, None, None

def compute_and_save(session_txns):
    all_txns = load_history_excel()
    all_txns.extend(session_txns)
    save_history_excel(all_txns)
    last = session_txns[-1]
    last_txn = {k: last[k] for k in ("date", "time", "merchant", "amount", "location")}
    total_txns   = len(all_txns)
    total_amount = sum(float(t["amount"]) for t in all_txns)
    avg_amount   = round(total_amount / total_txns, 2)

    def parse_date(val):
        if isinstance(val, datetime): return val
        return datetime.strptime(str(val), "%Y-%m-%d")

    dates     = [parse_date(t["date"]) for t in all_txns]
    earliest  = min(dates)
    latest    = max(dates)
    day_span  = max((latest - earliest).days + 1, 1)
    week_span = max(day_span / 7, 1)
    stats = {
        "total_txns":         total_txns,
        "avg_amount_per_txn": avg_amount,
        "avg_txns_per_day":   round(total_txns / day_span, 2),
        "avg_txns_per_week":  round(total_txns / week_span, 2),
    }
    merch_data = {}
    for t in all_txns:
        key = str(t["merchant"]).lower()
        if key not in merch_data:
            merch_data[key] = {"merchant": t["merchant"], "total": 0.0, "count": 0}
        merch_data[key]["total"] = round(merch_data[key]["total"] + float(t["amount"]), 2)
        merch_data[key]["count"] += 1
    merchant_averages = {
        key: {
            "merchant":   v["merchant"],
            "avg_amount": round(v["total"] / v["count"], 2),
            "count":      v["count"],
        }
        for key, v in merch_data.items()
    }
    save_summary_excel(last_txn, stats, merchant_averages)
    return all_txns, stats, merchant_averages

# ── Routes ────────────────────────────────────────────────────────────────────

@app.route("/api/user", methods=["GET"])
def get_user():
    return jsonify({"name": USER["name"]})

@app.route("/api/balance", methods=["POST"])
def get_balance():
    data = request.json or {}
    if data.get("pin") != USER["pin"]:
        return jsonify({"error": "Incorrect PIN"}), 401
    return jsonify({"balance": USER["balance"], "name": USER["name"]})

@app.route("/api/deposit", methods=["POST"])
def deposit():
    data = request.json or {}
    if data.get("pin") != USER["pin"]:
        return jsonify({"error": "Incorrect PIN"}), 401
    amount = float(data.get("amount", 0))
    if amount <= 0:
        return jsonify({"error": "Amount must be greater than zero"}), 400
    USER["balance"] = round(USER["balance"] + amount, 2)
    return jsonify({"balance": USER["balance"], "deposited": amount})

@app.route("/api/transaction", methods=["POST"])
def transaction():
    # Record the moment server receives the request (in ms)
    server_received_ts = time.time() * 1000

    data = request.json or {}
    if data.get("pin") != USER["pin"]:
        return jsonify({"error": "Incorrect PIN"}), 401

    amount   = float(data.get("amount", 0))
    merchant = str(data.get("merchant", "")).strip()
    category = data.get("category", "Others")
    txn_type = data.get("txn_type", "UPI")

    if amount <= 0:
        return jsonify({"error": "Amount must be greater than zero"}), 400
    if not merchant:
        return jsonify({"error": "Merchant name cannot be empty"}), 400
    if amount > USER["balance"]:
        return jsonify({"error": "Insufficient balance"}), 400

    # ── Timing from frontend timestamps ──────────────────────────────────────
    # session_start_ts → when user first opened the app (ms, sent by React)
    # txn_start_ts     → when user clicked the Pay button (ms, sent by React)
    session_start_ts = data.get("session_start_ts")
    txn_start_ts     = data.get("txn_start_ts")

    if session_start_ts:
        session_elapsed_s = (server_received_ts - float(session_start_ts)) / 1000
        session_time_str  = format_duration(session_elapsed_s)   # e.g. "4m 32.10s"
    else:
        session_time_str  = "Unknown"

    if txn_start_ts:
        # How many ms passed from user clicking Pay to server receiving it
        txn_duration_ms = round(server_received_ts - float(txn_start_ts), 2)
    else:
        txn_duration_ms = None
    # ─────────────────────────────────────────────────────────────────────────

    # Geo-locate
    try:
        from dotenv import load_dotenv
        load_dotenv()
        token    = os.getenv("IPINFO_TOKEN", "")
        res      = req_lib.get(f"https://ipinfo.io/json?token={token}", timeout=5)
        geo      = res.json()
        location = f"{geo.get('city')}, {geo.get('region')}, {geo.get('country')}"
    except Exception:
        location = "Unknown"

    now = datetime.now()
    txn = {
        "date":            now.strftime("%Y-%m-%d"),
        "time":            now.strftime("%H:%M:%S"),
        "merchant":        merchant,
        "category":        category,
        "txn_type":        txn_type,
        "amount":          amount,
        "net_debit":       amount,
        "location":        location,
        "session_time":    session_time_str,   # e.g. "4m 32.10s"
        "txn_duration_ms": txn_duration_ms,    # e.g. 2340.5 ms
    }

    USER["balance"] = round(USER["balance"] - amount, 2)

    return jsonify({
        "transaction":     txn,
        "balance":         USER["balance"],
        "session_time":    session_time_str,
        "txn_duration_ms": txn_duration_ms,
    })

@app.route("/api/history", methods=["GET"])
def history():
    return jsonify({"transactions": load_history_excel()})

@app.route("/api/summary", methods=["GET"])
def summary():
    last, stats, merch = load_summary_excel()
    if not last:
        return jsonify({"available": False})
    return jsonify({"available": True, "last": last, "stats": stats, "merchants": merch})

@app.route("/api/save", methods=["POST"])
def save():
    data         = request.json or {}
    session_txns = data.get("transactions", [])
    if not session_txns:
        return jsonify({"message": "No transactions to save"})
    all_txns, stats, merch = compute_and_save(session_txns)
    return jsonify({"message": "Saved successfully", "total": len(all_txns), "stats": stats})

if __name__ == "__main__":
    app.run(debug=True, port=5000)